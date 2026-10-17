"""Word-for-word check: every fragment of the locked slide copy (locked_copy.md) must appear in the deck.
Lines are split on table pipes, '; ', ' — ', ': ' and sentence ends, since the deck sets lists,
timeline headers and table cells as separate paragraphs or cells. Case-insensitive (chips use title case)."""
import re, sys
from pptx import Presentation
deck, src = sys.argv[1], sys.argv[2]
norm = lambda t: re.sub(r'\s+', ' ', t).strip().lower()
parts = []
for sl in Presentation(deck).slides:
    for sh in sl.shapes:
        if sh.has_text_frame: parts.append(sh.text_frame.text)
        if sh.has_table: parts += [c.text_frame.text for r in sh.table.rows for c in r.cells]
text = norm(' '.join(parts))
frags = []
for line in open(src, encoding='utf-8'):
    for f in re.split(r'\s\|\s|;\s|\s—\s|:\s|(?<=[.?])\s', line.strip()):
        f = f.strip(' .;:')
        if f and f != '○': frags.append(f)
missing = [f for f in frags if norm(f) not in text]
print(f'{len(frags) - len(missing)}/{len(frags)} locked fragments found verbatim')
for m in missing: print('MISSING:', m)
